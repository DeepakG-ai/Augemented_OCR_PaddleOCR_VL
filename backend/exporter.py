"""
exporter.py -- Build Excel and CSV exports for outbound delivery.
"""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_DATA_FONT = Font(name="Calibri", size=10)
_META_LABEL_FONT = Font(name="Calibri", bold=True, size=10)


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (list, dict)):
        return str(value)
    return str(value)


def _header_pairs(contract: dict[str, Any]) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for key, value in (contract.get("header") or {}).items():
        pairs.append((str(key), _stringify(value)))
    return pairs


def _line_items_table(contract: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    line_items = contract.get("line_items") or []
    columns: list[str] = []

    for item in line_items:
        if isinstance(item, dict):
            for key in item.keys():
                key_str = str(key)
                if key_str not in columns:
                    columns.append(key_str)

    if not columns:
        columns = ["line_items"]

    rows: list[list[Any]] = []
    if not line_items:
        return columns, rows

    for item in line_items:
        if isinstance(item, dict):
            rows.append([item.get(col) for col in columns])
        else:
            rows.append([_stringify(item)])
    return columns, rows


def _style_header_row(ws, row_number: int, col_count: int) -> None:
    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=row_number, column=col_idx)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER


def _style_data_rows(ws, start_row: int, col_count: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=col_count):
        for cell in row:
            cell.border = _THIN_BORDER
            cell.font = _DATA_FONT


def _auto_width(ws, min_width: int = 10, max_width: int = 40) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            value_len = len(_stringify(cell.value))
            if value_len > max_len:
                max_len = value_len
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


def build_excel_bytes(contract: dict[str, Any]) -> bytes:
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Field", "Value"])
    _style_header_row(ws_summary, 1, 2)

    for key, value in _header_pairs(contract):
        ws_summary.append([key, value])
        row_num = ws_summary.max_row
        ws_summary.cell(row=row_num, column=1).font = _META_LABEL_FONT
        ws_summary.cell(row=row_num, column=2).font = _DATA_FONT
        ws_summary.cell(row=row_num, column=1).border = _THIN_BORDER
        ws_summary.cell(row=row_num, column=2).border = _THIN_BORDER

    ws_summary.append([])
    metadata = [
        ("Extraction ID", contract.get("extraction_id")),
        ("Vendor ID", contract.get("vendor_id")),
        ("Vendor Name", contract.get("vendor_name")),
        ("Filename", contract.get("filename")),
        ("Canonical Source", (contract.get("review") or {}).get("canonical_source")),
        ("Reviewed At", (contract.get("review") or {}).get("reviewed_at")),
        ("Reason Code", (contract.get("review") or {}).get("reason_code")),
        ("Status", (contract.get("source") or {}).get("status")),
        ("Format Type", (contract.get("source") or {}).get("format_type")),
        ("Total Pages", (contract.get("source") or {}).get("total_pages")),
    ]
    for label, value in metadata:
        ws_summary.append([label, _stringify(value)])
        row_num = ws_summary.max_row
        ws_summary.cell(row=row_num, column=1).font = _META_LABEL_FONT
        ws_summary.cell(row=row_num, column=2).font = _DATA_FONT
        ws_summary.cell(row=row_num, column=1).border = _THIN_BORDER
        ws_summary.cell(row=row_num, column=2).border = _THIN_BORDER

    _auto_width(ws_summary)

    ws_items = wb.create_sheet("LineItems")
    columns, rows = _line_items_table(contract)
    ws_items.append(columns)
    _style_header_row(ws_items, 1, len(columns))

    for row_data in rows:
        ws_items.append(row_data)

    _style_data_rows(ws_items, 2, len(columns))
    _auto_width(ws_items)

    payload = BytesIO()
    wb.save(payload)
    return payload.getvalue()


def build_csv_bytes(contract: dict[str, Any]) -> bytes:
    header_pairs = _header_pairs(contract)
    header_keys = [key for key, _ in header_pairs]
    header_values = [value for _, value in header_pairs]

    columns, rows = _line_items_table(contract)
    all_columns = header_keys + columns

    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(all_columns)

    if rows:
        for row_data in rows:
            writer.writerow(header_values + [_stringify(value) for value in row_data])
    else:
        writer.writerow(header_values + [""] * len(columns))

    return buf.getvalue().encode("utf-8-sig")


__all__ = ["build_excel_bytes", "build_csv_bytes"]
